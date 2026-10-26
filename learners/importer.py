"""Bulk registration from the official Excel template (baseline §4).

Workflow: download template → upload → validate every row (formats, classes,
likely duplicates) → preview → confirm → create the valid rows in one
transaction → summary and error report. Uploaded rows are kept as plain
strings so the preview and the confirm step validate exactly the same data.
"""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from django.contrib.postgres.search import TrigramSimilarity
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.db.models.functions import Lower
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from core import audit
from core.phone import normalise_phone
from school.models import AcademicYear, ClassGroup

from . import services
from .duplicates import THRESHOLD, find_likely_duplicates
from .models import Guardian

COLUMNS = [
    # key, heading, required
    ("first_name", "First name", True),
    ("last_name", "Last name", True),
    ("other_names", "Other names", False),
    ("date_of_birth", "Date of birth (YYYY-MM-DD)", False),
    ("gender", "Gender (F/M)", False),
    ("class", "Class (e.g. 5A, F2B)", True),
    ("start_date", "Enrolment date (YYYY-MM-DD)", True),
    ("previous_school", "Previous school", False),
    ("guardian_name", "Guardian full name", True),
    ("relationship", "Relationship", False),
    ("guardian_phone", "Guardian mobile", True),
    ("guardian_email", "Guardian email", False),
]


def _heading_key(text):
    """'Class (e.g. 5A, F2B)' → 'class': the bracketed hint is optional when matching."""
    return re.sub(r"\s*\(.*$", "", text).strip().lower()


HEADINGS = {_heading_key(h): k for k, h, _ in COLUMNS} | {k: k for k, _, _ in COLUMNS}
GENDERS = {"F": "F", "FEMALE": "F", "M": "M", "MALE": "M"}
MAX_ROWS = 2000


class ImportFileError(Exception):
    pass


@dataclass
class Row:
    line: int
    raw: dict
    errors: list = field(default_factory=list)
    duplicates: list = field(default_factory=list)  # [(learner_number, full_name, why)]
    guardian_note: str = ""
    clean: dict = field(default_factory=dict)

    @property
    def ok(self):
        return not self.errors

    @property
    def name(self):
        return f"{self.raw.get('first_name', '')} {self.raw.get('last_name', '')}".strip()


@dataclass
class Preview:
    rows: list
    year: AcademicYear | None

    @property
    def valid(self):
        return [r for r in self.rows if r.ok]

    @property
    def invalid(self):
        return [r for r in self.rows if not r.ok]

    @property
    def flagged(self):
        return [r for r in self.valid if r.duplicates]


# Template and file reading ------------------------------------------------------------------

def build_template(year=None):
    year = year or AcademicYear.current()
    wb = Workbook()
    ws = wb.active
    ws.title = "Learners"
    ws.append([h for _, h, _ in COLUMNS])
    for i, (_, _, required) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        if required:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
        ws.column_dimensions[cell.column_letter].width = max(14, len(cell.value) + 2)
    help_ws = wb.create_sheet("Instructions")
    help_ws.append(["Fill one row per learner on the Learners sheet. Yellow columns are required."])
    help_ws.append(["Dates: YYYY-MM-DD, e.g. 2016-03-14. Gender: F or M. Mobile: Zambian number, e.g. 0971234567."])
    help_ws.append(["Learner numbers are created by the system; do not add a column for them."])
    help_ws.append([])
    help_ws.append([f"Classes in academic year {year}:" if year else "No academic year is current yet."])
    if year:
        for c in year.classes.filter(status="active").select_related("grade"):
            help_ws.append([c.name, c.grade.name])
    help_ws.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_upload(uploaded):
    """Return [(line, {key: text})] from an .xlsx or .csv upload."""
    name = (getattr(uploaded, "name", "") or "").lower()
    data = uploaded.read()
    if name.endswith(".csv"):
        table = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    elif name.endswith(".xlsx"):
        try:
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:  # openpyxl raises several types for a bad file
            raise ImportFileError("This file could not be read as an Excel workbook.") from exc
        table = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
    else:
        raise ImportFileError("Upload the Excel template (.xlsx) or a .csv file.")
    if not table:
        raise ImportFileError("The file is empty.")
    keys = [HEADINGS.get(_heading_key(_cell(h))) for h in table[0]]
    missing = [h for k, h, req in COLUMNS if req and k not in keys]
    if missing:
        raise ImportFileError(f"Missing column(s): {', '.join(missing)}. Use the official template.")
    rows = []
    for line, values in enumerate(table[1:], start=2):
        raw = {k: _cell(v) for k, v in zip(keys, values) if k}
        if any(raw.values()):
            rows.append((line, raw))
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"At most {MAX_ROWS} learners per file; split the file.")
    if not rows:
        raise ImportFileError("No learners found in the file.")
    return rows


# Validation ---------------------------------------------------------------------------------

def _parse_date(text):
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError


def validate(raw_rows, school, year=None):
    year = year or AcademicYear.current(school)
    classes = {}
    if year:
        classes = {c.name.upper(): c for c in ClassGroup.objects.filter(academic_year=year, status="active").select_related("grade", "academic_year")}
    rows = [Row(line, raw) for line, raw in raw_rows]
    seen_people = {}
    file_guardians = {}

    for row in rows:
        r, e = row.raw, row.errors
        for key, heading, required in COLUMNS:
            if required and not r.get(key):
                e.append(f"{heading} is required")
        clean = row.clean
        for key in ("first_name", "last_name", "other_names", "previous_school", "guardian_name", "relationship"):
            clean[key] = r.get(key, "")
        for key in ("date_of_birth", "start_date"):
            clean[key] = None
            if r.get(key):
                try:
                    clean[key] = _parse_date(r[key])
                except ValueError:
                    e.append(f"{dict((k, h) for k, h, _ in COLUMNS)[key]}: '{r[key]}' is not a date (use YYYY-MM-DD)")
        if clean["date_of_birth"] and clean["date_of_birth"] > date.today():
            e.append("Date of birth is in the future")
        gender = r.get("gender", "")
        clean["gender"] = GENDERS.get(gender.upper(), "") if gender else ""
        if gender and not clean["gender"]:
            e.append("Gender must be F or M")
        clean["class_group"] = None
        if r.get("class"):
            if year is None:
                e.append("No academic year is current; ask an administrator to set one")
            else:
                cg = classes.get(r["class"].upper().replace(" ", "").replace("GRADE", "").replace("FORM", "F"))
                if cg is None:
                    e.append(f"Class '{r['class']}' does not exist in {year}")
                else:
                    clean["class_group"] = cg
                    if clean["start_date"] and not year.contains(clean["start_date"]):
                        e.append(f"Enrolment date must fall inside {year} ({year.start_date} to {year.end_date})")
        clean["guardian_phone"] = ""
        if r.get("guardian_phone"):
            try:
                clean["guardian_phone"] = normalise_phone(r["guardian_phone"])
            except ValidationError:
                e.append(f"Guardian mobile '{r['guardian_phone']}' is not a valid Zambian mobile number")
        clean["guardian_email"] = r.get("guardian_email", "")
        if clean["guardian_email"]:
            try:
                validate_email(clean["guardian_email"])
            except ValidationError:
                e.append(f"Guardian email '{clean['guardian_email']}' is not valid")
        if e:
            continue

        # Likely duplicates: against existing learners, then within the file.
        for learner, why in find_likely_duplicates(
            school, clean["first_name"], clean["last_name"], clean["date_of_birth"], clean["guardian_phone"]
        ):
            row.duplicates.append((learner.learner_number, learner.full_name, why))
        person = (clean["first_name"].lower(), clean["last_name"].lower(), clean["date_of_birth"])
        if person in seen_people:
            row.duplicates.append(("", row.name, f"same name and date of birth as line {seen_people[person]} of this file"))
        else:
            seen_people[person] = row.line

        # Guardian: link an existing one only when phone matches and the name is similar (D3); say so in the preview.
        existing = (
            Guardian.objects.filter(school=school, contacts__phone=clean["guardian_phone"], contacts__is_current=True)
            .annotate(similarity=TrigramSimilarity(Lower("full_name"), clean["guardian_name"].lower()))
            .filter(similarity__gte=THRESHOLD)
            .order_by("-similarity")
            .first()
        )
        file_key = (clean["guardian_phone"], clean["guardian_name"].lower())
        if existing:
            clean["guardian_id"] = existing.pk
            row.guardian_note = f"Links to existing guardian {existing.full_name}"
        elif file_key in file_guardians:
            clean["guardian_line"] = file_guardians[file_key]
            row.guardian_note = f"Same new guardian as line {file_guardians[file_key]}"
        else:
            file_guardians[file_key] = row.line
            others = Guardian.objects.filter(school=school, contacts__phone=clean["guardian_phone"], contacts__is_current=True)
            row.guardian_note = "New guardian" + (
                f" (this number is also used by {', '.join(g.full_name for g in others[:3])})" if others.exists() else ""
            )
    return Preview(rows, year)


# Confirm ------------------------------------------------------------------------------------

@dataclass
class ImportResult:
    created: list = field(default_factory=list)  # [(line, learner)]
    skipped_duplicates: list = field(default_factory=list)  # [Row]
    invalid: list = field(default_factory=list)  # [Row]


def confirm(raw_rows, school, actor, *, override_reason="", year=None):
    """Create every valid row in one transaction. Rows flagged as likely duplicates need an override reason."""
    if not actor.has_perm("learners.import_learners"):
        raise PermissionDenied("You may not import learners.")
    override_reason = override_reason.strip()
    if override_reason and not actor.has_perm("learners.override_duplicate"):
        raise PermissionDenied("You may not import likely duplicates (D5).")
    preview = validate(raw_rows, school, year)
    result = ImportResult(invalid=preview.invalid)
    new_guardians = {}
    with transaction.atomic():
        for row in preview.valid:
            if row.duplicates and not override_reason:
                result.skipped_duplicates.append(row)
                continue
            c = row.clean
            guardian = None
            if c.get("guardian_id"):
                guardian = Guardian.objects.get(pk=c["guardian_id"])
            elif c.get("guardian_line"):
                guardian = new_guardians.get(c["guardian_line"])
            reg = services.Registration(
                first_name=c["first_name"], last_name=c["last_name"], other_names=c["other_names"],
                date_of_birth=c["date_of_birth"], gender=c["gender"], previous_school=c["previous_school"],
                class_group=c["class_group"], start_date=c["start_date"],
                guardian=services.GuardianInput(c["guardian_name"], c["guardian_phone"], c["relationship"],
                                                c["guardian_email"], existing=guardian),
            )
            learner = services.register_learner(reg, actor, school=school, check_duplicates=False)
            if row.duplicates:
                audit.record("duplicate.override", learner, actor=actor, reason=override_reason,
                             changes={"import_line": [None, row.line],
                                      "matched": [None, [d[0] or d[2] for d in row.duplicates]]})
            if guardian is None:
                new_guardians[row.line] = learner.guardian_links.get().guardian
            result.created.append((row.line, learner))
        audit.record("learner_import.confirm", None, school=school, actor=actor, object_type="learner_import",
                     reason=override_reason,
                     changes={"rows": [None, len(preview.rows)], "created": [None, len(result.created)],
                              "invalid": [None, len(result.invalid)],
                              "skipped_duplicates": [None, len(result.skipped_duplicates)],
                              "learner_numbers": [None, [l.learner_number for _, l in result.created]]})
    return result


def error_report(rows):
    """The rows that were not imported, with a Problems column, ready to fix and re-upload."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Learners"
    ws.append([h for _, h, _ in COLUMNS] + ["Problems (delete this column before re-uploading)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        problems = row.errors or [f"Likely duplicate: {d[1]} {d[0]} ({d[2]})" for d in row.duplicates]
        ws.append([row.raw.get(k, "") for k, _, _ in COLUMNS] + ["; ".join(problems)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

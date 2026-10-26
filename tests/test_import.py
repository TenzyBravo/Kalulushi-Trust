import datetime
import io

from django.core.exceptions import PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from accounts import roles
from core.models import AuditLog, School
from learners import importer
from learners.models import Guardian, Learner

from . import factories as f

HEAD = [h for _, h, _ in importer.COLUMNS]


def xlsx(rows, headings=HEAD):
    wb = Workbook()
    ws = wb.active
    ws.append(headings)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("learners.xlsx", buf.getvalue())


def row(first, last, dob, cls, guardian, phone, gender="F", start="2026-01-20"):
    return [first, last, "", dob, gender, cls, start, "", guardian, "Mother", phone, ""]


GOOD = [
    row("Chilufya", "Mwansa", "2016-03-14", "4A", "Mary Mwansa", "0971234567"),
    row("Mwila", "Mwansa", "2018-07-02", "2A", "Mary Mwansa", "+260 97 123 4567", gender="M"),
    row("Natasha", "Banda", datetime.date(2012, 11, 25), "Form 1A", "Peter Banda", "0961112233"),
]


class ImportTests(TestCase):
    def setUp(self):
        self.admin = f.admin_user()
        f.year()
        for code in ("2", "4", "F1"):
            f.make_class(code)
        self.school = School.current()

    def rows(self, data):
        return importer.read_upload(xlsx(data))

    def test_template_has_the_columns_and_classes(self):
        wb = load_workbook(io.BytesIO(importer.build_template()))
        self.assertEqual([c.value for c in wb["Learners"][1]], HEAD)
        self.assertIn("F1A", [r[0].value for r in wb["Instructions"].iter_rows()])

    def test_preview_finds_problems_without_saving(self):
        bad = GOOD + [row("Bwalya", "", "14/14/2015", "9Z", "Ann Zulu", "123", gender="X")]
        preview = importer.validate(self.rows(bad), self.school)
        self.assertEqual(len(preview.valid), 3)
        errors = preview.invalid[0].errors
        self.assertIn("Last name is required", errors)
        self.assertTrue(any("not a date" in e for e in errors))
        self.assertTrue(any("Class '9Z'" in e for e in errors))
        self.assertTrue(any("valid Zambian mobile" in e for e in errors))
        self.assertIn("Gender must be F or M", errors)
        self.assertFalse(Learner.objects.exists())

    def test_confirm_creates_valid_rows_and_shares_the_guardian(self):
        result = importer.confirm(self.rows(GOOD), self.school, self.admin)
        self.assertEqual(len(result.created), 3)
        self.assertEqual(Guardian.objects.count(), 2)  # siblings share Mary
        mary = Guardian.objects.get(full_name="Mary Mwansa")
        self.assertEqual(mary.learners.count(), 2)
        self.assertEqual(Learner.objects.get(first_name="Natasha").current_enrollment().class_group.name, "F1A")
        summary = AuditLog.objects.get(action="learner_import.confirm")
        self.assertEqual(summary.changes["created"], [None, 3])

    def test_reimporting_flags_duplicates_and_skips_them(self):
        importer.confirm(self.rows(GOOD), self.school, self.admin)
        preview = importer.validate(self.rows(GOOD), self.school)
        self.assertEqual(len(preview.flagged), 3)
        self.assertIn("Links to existing guardian Mary Mwansa", preview.rows[0].guardian_note)
        result = importer.confirm(self.rows(GOOD), self.school, self.admin)
        self.assertEqual((len(result.created), len(result.skipped_duplicates)), (0, 3))
        self.assertEqual(Learner.objects.count(), 3)

    def test_override_imports_flagged_rows_with_reason(self):
        importer.confirm(self.rows(GOOD[:1]), self.school, self.admin)
        result = importer.confirm(self.rows(GOOD[:1]), self.school, self.admin, override_reason="Twins")
        self.assertEqual(len(result.created), 1)
        self.assertTrue(AuditLog.objects.filter(action="duplicate.override", reason="Twins").exists())

    def test_teachers_cannot_import(self):
        with self.assertRaises(PermissionDenied):
            importer.confirm(self.rows(GOOD), self.school, f.make_user("t", roles.TEACHER))

    def test_wrong_file_is_rejected(self):
        with self.assertRaises(importer.ImportFileError):
            importer.read_upload(xlsx([["x"]], headings=["Name"]))
        with self.assertRaises(importer.ImportFileError):
            importer.read_upload(SimpleUploadedFile("learners.pdf", b"%PDF"))

    def test_error_report_lists_rows_with_problems(self):
        preview = importer.validate(self.rows([row("A", "", "", "4A", "B", "1")]), self.school)
        wb = load_workbook(io.BytesIO(importer.error_report(preview.invalid)))
        last = [c.value for c in wb.active[2]][-1]
        self.assertIn("Last name is required", last)


class ImportCommandTests(TestCase):
    def test_command_checks_then_imports(self):
        from io import StringIO

        from django.conf import settings
        from django.core.management import call_command

        f.admin_user()
        f.year()
        for code in ("2", "4", "F1"):
            f.make_class(code)
        path = str(settings.BASE_DIR / "samples" / "learners.csv")
        out = StringIO()
        call_command("import_learners", path, as_user="sysadmin", stdout=out, stderr=StringIO())
        self.assertIn("3 valid", out.getvalue())
        self.assertFalse(Learner.objects.exists())
        call_command("import_learners", path, as_user="sysadmin", confirm=True, stdout=out, stderr=StringIO())
        self.assertEqual(Learner.objects.count(), 3)
